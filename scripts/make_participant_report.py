#!/usr/bin/env python
"""
Individual single-participant report (.docx + figure).

Shows one participant's Performance (Cognition), Mood (Emotion) and Awareness
(permutation entropy, plus the Lempel-Ziv companion) across the protocol,
against their OWN baseline, with the rest of the cohort as grey context.
With n=1 there are no significance tests -- the report is descriptive, flags
how little usable data there is, and says so plainly.

Usage:
    .venv/bin/python scripts/make_participant_report.py --subject Sub_53 \
        --label "Gerard Walsh" --outdir output/participant_reports
"""
from __future__ import annotations

import argparse
import datetime as dt
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

MODULE_ORDER = [
    "PSYNAUTICS_BASELINE_EEG_1", "PSYNAUTICS_BASELINE_EEG_2", "PSYNAUTICS_BASELINE_EEG_3",
    "PSYNAUTICS_BASELINE_EEG_4", "PSYNAUTICS_BASELINE_EEG_5", "PSYNAUTICS_BASELINE_EEG_TDAY",
    "PSYNAUTICS_EEG_TREATMENT", "PSYNAUTICS_POST_EEG_1", "PSYNAUTICS_POST_EEG_2",
    "PSYNAUTICS_POST_EEG_3", "PSYNAUTICS_POST_EEG_4",
]
SHORT = ["Base 1", "Base 2", "Base 3", "Base 4", "Base 5", "Base\n(Tx day)", "Treatment", "Post 1", "Post 2", "Post 3", "Post 4"]
BASELINE = {m for m in MODULE_ORDER if "BASELINE" in m}
POST = {m for m in MODULE_ORDER if "POST" in m}
TREATMENT = "PSYNAUTICS_EEG_TREATMENT"

# csv column, display name, log10?, color, unit text
MARKERS = [
    ("cognition", "Performance (AF8 broadband 1-44Hz)", True, "#4C72B0", "uV^2/Hz"),
    ("emotion", "Mood (AF7 alpha 7.5-13Hz)", True, "#C44E52", "uV^2/Hz"),
    ("awareness", "Awareness (permutation entropy, 4 ch)", False, "#55A868", ""),
    ("awareness_lzc", "Awareness companion (Lempel-Ziv, 4 ch)", False, "#8172B2", ""),
]
GREY = RGBColor(0x55, 0x55, 0x55)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), fill)
    tcPr.append(s)


def para(doc, text, bold=False, italic=False, color=None, size=None):
    if text.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        text = text[2:]
    else:
        p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(8)
        shade(c, "D9E2F3")
        c.width = Inches(widths[i])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(8)
            if ri % 2:
                shade(cells[i], "F2F2F2")
            cells[i].width = Inches(widths[i])
    return t


def val(x, log):
    return np.log10(x) if log else x


def back(x, log):
    return 10 ** x if log else x


def fmt(x, log):
    if x is None or pd.isna(x):
        return "n/a"
    return f"{x:,.1f}" if log else f"{x:.3f}"


def load_all(marker_csv: Path, qc_csv: Path) -> pd.DataFrame:
    m = pd.read_csv(marker_csv)
    q = pd.read_csv(qc_csv)[["session_hash", "datetime", "duration_s", "pct_annotated_bad"]]
    d = m.merge(q, on="session_hash", how="left")
    d["datetime"] = pd.to_datetime(d["datetime"], utc=True)
    d["step"] = d["module_id"].map({m: i for i, m in enumerate(MODULE_ORDER)})
    return d


def make_figure(me: pd.DataFrame, others: pd.DataFrame, out: Path, label: str) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(13, 8.5))
    for ax, (col, name, log, color, _u) in zip(axes.ravel(), MARKERS):
        o = others.dropna(subset=[col]).copy()
        o["v"] = val(o[col], log)
        g = o.groupby("step")["v"]
        steps = sorted(g.groups)
        lo = [g.get_group(s).quantile(0.25) for s in steps]
        hi = [g.get_group(s).quantile(0.75) for s in steps]
        med = [g.get_group(s).median() for s in steps]
        ax.fill_between(steps, back(np.array(lo), log), back(np.array(hi), log), color="#CCCCCC", alpha=0.6, label="Rest of cohort (IQR)")
        ax.plot(steps, back(np.array(med), log), color="#888888", linewidth=1, label="Cohort median")

        mine = me.dropna(subset=[col]).copy()
        mine["v"] = val(mine[col], log)
        base = mine[mine["module_id"].isin(BASELINE)]
        if len(base):
            ax.axhspan(back(base["v"].min(), log), back(base["v"].max(), log), xmin=0, xmax=1, color=color, alpha=0.12)
            ax.axhline(back(base["v"].mean(), log), color=color, linestyle=":", linewidth=1, label="Own baseline mean / range")
        rng = np.random.default_rng(0)
        xs = mine["step"].to_numpy() + rng.uniform(-0.12, 0.12, len(mine))
        ax.scatter(xs, mine[col], color=color, s=45, zorder=5, edgecolor="k", linewidth=0.5, label=f"{label}")
        if log:
            ax.set_yscale("log")
        ax.axvline(MODULE_ORDER.index(TREATMENT) - 0.5, color="gray", linestyle="--", linewidth=0.8)
        ax.axvline(MODULE_ORDER.index(TREATMENT) + 0.5, color="gray", linestyle="--", linewidth=0.8)
        ax.set_xticks(range(len(SHORT)))
        ax.set_xticklabels(SHORT, fontsize=7)
        ax.set_ylabel(_u if _u else "normalised, 0-1", fontsize=8)
        n_ok = int(me[col].notna().sum())
        ax.set_title(f"{name}\nusable in {n_ok} of {len(me)} recordings", fontsize=10)
        if ax is axes[0, 0]:
            ax.legend(fontsize=7, loc="best")
    fig.suptitle(f"{label}: markers across the protocol vs. own baseline and the rest of the cohort", fontsize=13)
    fig.tight_layout()
    fig.savefig(out, dpi=150)
    plt.close(fig)


def summarise(me: pd.DataFrame, col: str, log: bool) -> dict:
    s = me.dropna(subset=[col]).copy()
    s["v"] = val(s[col], log)
    base = s[s["module_id"].isin(BASELINE)]
    post = s[s["module_id"].isin(POST)]
    tx = s[s["module_id"] == TREATMENT]
    out = {"n_base": len(base), "n_post": len(post), "n_tx": len(tx)}
    out["base"] = back(base["v"].mean(), log) if len(base) else None
    out["post"] = back(post["v"].mean(), log) if len(post) else None
    out["tx"] = back(tx["v"].mean(), log) if len(tx) else None
    out["base_rng"] = (back(base["v"].min(), log), back(base["v"].max(), log)) if len(base) else None
    if out["base"] is not None and out["post"] is not None:
        out["change"] = (out["post"] / out["base"]) if log else (out["post"] - out["base"])
        lo, hi = out["base_rng"]
        out["inside"] = bool(lo <= out["post"] <= hi)
    return out


def build_doc(me, others, label, subject, fig_path, clin_rows, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(0.75)

    doc.add_heading(f"Individual Study Report: {label}", level=0)
    para(doc, f"Psynautics EEG pipeline · Muse S resting-state recordings · generated {dt.date.today():%d %b %Y}", italic=True, color=GREY)

    doc.add_heading("1. Summary", level=1)
    first, last = me["datetime"].min(), me["datetime"].max()
    para(doc, f"{len(me)} recordings between {first:%d %b %Y} and {last:%d %b %Y}, covering "
              f"{me['module_id'].nunique()} of the 11 protocol steps.")
    sums = {c: summarise(me, c, lg) for c, _n, lg, _c, _u in MARKERS}
    for col, name, log, _c, unit in MARKERS:
        s = sums[col]
        if s["n_base"] == 0 or s["n_post"] == 0:
            para(doc, f"- {name}: cannot compare baseline with post-treatment "
                      f"({s['n_base']} usable baseline, {s['n_post']} usable post-treatment recordings).")
            continue
        chg = f"x{s['change']:.2f}" if log else f"{s['change']:+.3f}"
        where = "inside" if s["inside"] else "outside"
        warn = " Too few baseline sessions to judge reliably." if s["n_base"] < 3 else ""
        lo_, hi_ = s["base_rng"]
        if log and lo_ > 0 and hi_ / lo_ > 10:
            warn += f" The baseline itself is very variable (highest session is {hi_ / lo_:.0f}x the lowest), so a wide 'inside baseline' range says little."
        para(doc, f"- {name}: baseline mean {fmt(s['base'], log)} (n={s['n_base']}), post-treatment mean "
                  f"{fmt(s['post'], log)} (n={s['n_post']}), change {chg}; post-treatment is {where} the "
                  f"participant's own baseline range.{warn}")
    para(doc, "These are descriptive observations on one person. No significance tests are possible with n=1, and a "
              "value sitting outside the baseline range may simply reflect recording noise (see Data quality).",
         italic=True, color=GREY)

    doc.add_heading("2. Markers across the protocol", level=1)
    doc.add_picture(str(fig_path), width=Inches(6.9))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Coloured dots are this participant's recordings. The shaded colour band is their own baseline range "
              "(dotted line = baseline mean). Grey band and line = middle 50% and median of all other participants at "
              "each protocol step. Dashed vertical lines bracket the treatment-day session. Performance and Mood are on a "
              "log scale.", size=9)

    doc.add_heading("3. Every recording", level=1)
    rows = []
    for _, r in me.sort_values("datetime").iterrows():
        rows.append([
            SHORT[int(r["step"])].replace("\n", " "), f"{r['datetime']:%Y-%m-%d}", f"{r['duration_s'] / 60:.1f}",
            f"{r['n_epochs_used']:.0f}", f"{r['pct_annotated_bad']:.0f}%", (r["bad_channels"] if isinstance(r["bad_channels"], str) else "-"),
            fmt(r["cognition"], True), fmt(r["emotion"], True), fmt(r["awareness"], False), fmt(r["awareness_lzc"], False),
        ])
    table(doc, ["Step", "Date", "Min", "Usable\nepochs", "Time\nbad", "Bad channels", "Perf.", "Mood", "Aware.", "LZC"],
          rows, [0.7, 0.8, 0.4, 0.55, 0.5, 0.95, 0.75, 0.8, 0.6, 0.55])
    para(doc, "Perf. and Mood in uV^2/Hz; Awareness and LZC are 0-1. n/a = not computable because a required channel was "
              "flagged bad for that recording. 'Usable epochs' = 4-second windows left after removing dropout/artifact/muscle windows.", size=8, italic=True)

    doc.add_heading("4. Baseline vs. post-treatment", level=1)
    rows = []
    for col, name, log, _c, _u in MARKERS:
        s = sums[col]
        rng = f"{fmt(s['base_rng'][0], log)} - {fmt(s['base_rng'][1], log)}" if s["base_rng"] else "n/a"
        chg = "n/a" if "change" not in s else (f"x{s['change']:.2f}" if log else f"{s['change']:+.3f}")
        rows.append([name.split(" (")[0], f"{fmt(s['base'], log)} (n={s['n_base']})", rng,
                     f"{fmt(s['tx'], log)} (n={s['n_tx']})", f"{fmt(s['post'], log)} (n={s['n_post']})", chg])
    table(doc, ["Marker", "Baseline mean*", "Baseline range", "Treatment day", "Post-treatment mean", "Change vs baseline"],
          rows, [1.1, 1.2, 1.4, 1.1, 1.3, 0.9])

    para(doc, "*Performance and Mood means are geometric (averaged on a log scale), because the raw values span orders of magnitude.", size=8, italic=True)

    doc.add_heading("5. Data quality", level=1)
    bc = me["bad_channels"].fillna("")
    for ch in ["TP9", "AF7", "AF8", "TP10"]:
        k = int(bc.str.contains(ch).sum())
        para(doc, f"- {ch}: flagged bad in {k} of {len(me)} recordings.")
    para(doc, f"- Median share of each recording lost to dropout/artifact/muscle: {me['pct_annotated_bad'].median():.0f}% "
              f"(cohort median {others['pct_annotated_bad'].median():.0f}%).")
    para(doc, f"- Performance usable in {int(me['cognition'].notna().sum())}, Mood in {int(me['emotion'].notna().sum())}, "
              f"Awareness in {int(me['awareness'].notna().sum())} of {len(me)} recordings. Awareness needs all four channels clean at once, "
              "so it is the first to disappear when electrode contact is poor.")
    order_notes = []
    dup = me["module_id"].value_counts()
    for m_, k in dup[dup > 1].items():
        order_notes.append(f"{SHORT[MODULE_ORDER.index(m_)].replace(chr(10), ' ')} has {k} recordings")
    ordered = me.sort_values("datetime")["step"].to_numpy()
    if (np.diff(ordered) < 0).any():
        order_notes.append("recording dates are not in protocol order")
    missing = [SHORT[i].replace("\n", " ") for i in range(len(MODULE_ORDER)) if i not in set(me["step"].astype(int))]
    if missing:
        order_notes.append("no recording for: " + ", ".join(missing))
    if order_notes:
        para(doc, "- Protocol irregularities: " + "; ".join(order_notes) + ". Step labels come from the app; the pipeline uses them as-is.")

    doc.add_heading("6. Self-reported treatment details", level=1)
    if clin_rows:
        for r in clin_rows:
            para(doc, f"- Form dated {r['date']}: dose \"{r['Dose'] or 'not given'}\"; route \"{r['RouteOfAdmin'] or 'not given'}\"; "
                      f"clinic/referral \"{r['ClinicDetails'] or 'not given'}\"; additional \"{r['AdditionalInfo'] or 'none'}\".")
        para(doc, "Free-text, self-reported and unverified. Not used in any calculation above.", italic=True, color=GREY)
    else:
        para(doc, "No clinic-info form found for this participant.")

    doc.add_heading("7. How to read this", level=1)
    for t in [
        "Everything is descriptive: one participant cannot support statistical inference.",
        "Performance (AF8) and Mood (AF7) are frontal dry-electrode measures and are partly contaminated by facial/jaw muscle "
        "activity, which the pipeline only partly removes. A change in either may reflect muscle tension rather than brain activity.",
        "Awareness (permutation entropy) is highly stable within a person and moves very little, so small shifts are hard to interpret.",
        "When most recordings have a bad ear electrode (TP9/TP10), the 'grey cohort band' is a comparison against cleaner data than this person's.",
    ]:
        doc.add_paragraph(t, style="List Bullet")
    doc.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--subject", required=True, help="Subject code, e.g. Sub_53")
    ap.add_argument("--label", default=None, help="Name shown in the report (default: the subject code)")
    ap.add_argument("--marker-summary", type=Path, default=Path("output/batch_full/marker_summary.csv"))
    ap.add_argument("--qc-summary", type=Path, default=Path("output/batch_full/batch_qc_summary.csv"))
    ap.add_argument("--clinicinfo-dir", type=Path, default=Path("data/clinicinfo"))
    ap.add_argument("--participant-key", type=Path, default=Path("/Users/sonyaluchanskaya/Desktop/Psynautics/participant_key_CONFIDENTIAL.csv"))
    ap.add_argument("--outdir", type=Path, default=Path("output/participant_reports"))
    a = ap.parse_args()

    a.outdir.mkdir(parents=True, exist_ok=True)
    label = a.label or a.subject
    d = load_all(a.marker_summary, a.qc_summary)
    me, others = d[d["subject_code"] == a.subject], d[d["subject_code"] != a.subject]
    if me.empty:
        raise SystemExit(f"No recordings found for {a.subject}")

    key = pd.read_csv(a.participant_key, header=None, names=["s", "e"])
    emails = key.loc[key["s"] == a.subject, "e"].tolist()
    clin = []
    for f in sorted(a.clinicinfo_dir.glob("*.csv")):
        em, date = f.stem.rsplit("__", 1)
        if em in emails:
            x = pd.read_csv(f)
            row = dict(zip(x["program_question_id"], x["latest_answer"].fillna("").astype(str)))
            clin.append({"date": date, **{k: " ".join(row.get(k, "").split()) for k in ["Dose", "RouteOfAdmin", "ClinicDetails", "AdditionalInfo"]}})

    fig_path = a.outdir / f"{a.subject}_markers.png"
    make_figure(me, others, fig_path, label)
    out = a.outdir / f"{a.subject}_study_report.docx"
    build_doc(me, others, label, a.subject, fig_path, clin, out)
    print(f"Saved {out} and {fig_path}")


if __name__ == "__main__":
    main()
